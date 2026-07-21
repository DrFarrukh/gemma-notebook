import csv
import html
import io
import math
import os
import re
import struct
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from docx import Document
from odf import teletype
from odf.opendocument import load as load_odf
from pypdf import PdfReader

from . import db

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")
GENERATION_MODEL = os.getenv("GENERATION_MODEL", "gemma4-e4b-64k:latest")
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "nomic-embed-text:latest")
MAX_CHARS = 2_000_000
CHUNK_SIZE = 1800
CHUNK_OVERLAP = 250


def clean_text(text):
    text = text.replace("\x00", " ").replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(path: Path, kind: str):
    suffix = path.suffix.lower()
    pages = []
    scanned = False
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        for index, page in enumerate(reader.pages, 1):
            pages.append({"page": index, "text": clean_text(page.extract_text() or "")})
        total = sum(len(p["text"]) for p in pages)
        scanned = bool(pages and total < max(80, len(pages) * 30))
    elif suffix == ".docx":
        document = Document(str(path))
        text = "\n\n".join(p.text for p in document.paragraphs if p.text.strip())
        pages = [{"page": None, "text": clean_text(text)}]
    elif suffix == ".odt":
        document = load_odf(str(path))
        pages = [{"page": None, "text": clean_text(teletype.extractText(document.text))}]
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(raw, "html.parser")
            for tag in soup(["script", "style", "noscript"]):
                tag.decompose()
            raw = soup.get_text("\n")
        elif suffix == ".csv":
            parsed = csv.reader(io.StringIO(raw))
            raw = "\n".join(" | ".join(cell.strip() for cell in row) for row in parsed)
        pages = [{"page": None, "text": clean_text(raw)}]
    combined = "\n\n".join(p["text"] for p in pages)
    if len(combined) > MAX_CHARS:
        raise ValueError(f"Extracted text exceeds the {MAX_CHARS:,}-character source limit")
    if not combined and not scanned:
        raise ValueError("No readable text was found in this source")
    return pages, scanned


def split_pages(pages):
    chunks = []
    ordinal = 0
    for page in pages:
        text = page["text"]
        start = 0
        while start < len(text):
            end = min(start + CHUNK_SIZE, len(text))
            if end < len(text):
                boundary = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
                if boundary > start + CHUNK_SIZE // 2:
                    end = boundary + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append({"ordinal": ordinal, "page": page["page"], "text": piece})
                ordinal += 1
            if end >= len(text):
                break
            start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def pack_vector(values):
    return struct.pack(f"<{len(values)}f", *values)


def unpack_vector(blob):
    return struct.unpack(f"<{len(blob) // 4}f", blob)


def cosine(a, b):
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(x * x for x in b))
    return dot / (na * nb) if na and nb else 0.0


def embed_texts(texts):
    if not texts:
        return []
    with httpx.Client(timeout=180) as client:
        response = client.post(f"{OLLAMA_URL}/api/embed", json={"model": EMBEDDING_MODEL, "input": texts})
        response.raise_for_status()
        return response.json()["embeddings"]


def process_source(source_id):
    source = db.row("SELECT * FROM sources WHERE id=?", (source_id,))
    if not source:
        return
    db.execute("UPDATE sources SET status='processing', error=NULL, updated_at=? WHERE id=?", (db.now(), source_id))
    try:
        pages, scanned = extract(Path(source["path"]), source["kind"])
        if scanned:
            db.execute(
                "UPDATE sources SET status='error', scanned=1, error=?, page_count=?, updated_at=? WHERE id=?",
                ("This appears to be an image-only PDF. Local OCR is not enabled in v1.", len(pages), db.now(), source_id),
            )
            return
        chunks = split_pages(pages)
        vectors = []
        for offset in range(0, len(chunks), 24):
            vectors.extend(embed_texts([c["text"] for c in chunks[offset:offset + 24]]))
        with db.connection() as conn:
            conn.execute("DELETE FROM chunk_fts WHERE chunk_id IN (SELECT id FROM chunks WHERE source_id=?)", (source_id,))
            conn.execute("DELETE FROM chunks WHERE source_id=?", (source_id,))
            for chunk, vector in zip(chunks, vectors):
                chunk_id = db.uid()
                conn.execute(
                    "INSERT INTO chunks(id,source_id,notebook_id,ordinal,page,section,text,embedding) VALUES(?,?,?,?,?,?,?,?)",
                    (chunk_id, source_id, source["notebook_id"], chunk["ordinal"], chunk["page"], None, chunk["text"], pack_vector(vector)),
                )
                conn.execute("INSERT INTO chunk_fts(chunk_id,text) VALUES(?,?)", (chunk_id, chunk["text"]))
            count = sum(len(p["text"]) for p in pages)
            conn.execute(
                "UPDATE sources SET status='ready', scanned=0, error=NULL, char_count=?, page_count=?, updated_at=? WHERE id=?",
                (count, len(pages) if source["kind"] == "pdf" else None, db.now(), source_id),
            )
    except Exception as exc:
        db.execute("UPDATE sources SET status='error', error=?, updated_at=? WHERE id=?", (str(exc)[:1000], db.now(), source_id))


def _fts_query(query):
    tokens = re.findall(r"[\w]{2,}", query.lower(), re.UNICODE)[:12]
    return " OR ".join(f'"{token}"' for token in tokens)


def retrieve(notebook_id, query, source_ids=None, limit=12):
    params = [notebook_id]
    clause = "c.notebook_id=? AND s.status='ready' AND s.enabled=1"
    if source_ids:
        placeholders = ",".join("?" for _ in source_ids)
        clause += f" AND c.source_id IN ({placeholders})"
        params.extend(source_ids)
    candidates = db.rows(
        f"SELECT c.*,s.name AS source_name FROM chunks c JOIN sources s ON s.id=c.source_id WHERE {clause}", params
    )
    if not candidates:
        return []
    query_vector = embed_texts([query])[0]
    semantic = sorted(candidates, key=lambda c: cosine(query_vector, unpack_vector(c["embedding"])), reverse=True)
    semantic_rank = {item["id"]: rank for rank, item in enumerate(semantic[:80], 1)}
    lexical_rank = {}
    fts = _fts_query(query)
    if fts:
        fts_params = [fts, notebook_id]
        source_filter = ""
        if source_ids:
            source_filter = f" AND c.source_id IN ({','.join('?' for _ in source_ids)})"
            fts_params.extend(source_ids)
        try:
            matches = db.rows(
                "SELECT c.id FROM chunk_fts f JOIN chunks c ON c.id=f.chunk_id "
                "JOIN sources s ON s.id=c.source_id WHERE chunk_fts MATCH ? AND c.notebook_id=? "
                f"AND s.status='ready' AND s.enabled=1 {source_filter} ORDER BY bm25(chunk_fts) LIMIT 80",
                fts_params,
            )
            lexical_rank = {item["id"]: rank for rank, item in enumerate(matches, 1)}
        except Exception:
            lexical_rank = {}
    for item in candidates:
        item["score"] = 0.0
        if item["id"] in semantic_rank:
            item["score"] += 1 / (60 + semantic_rank[item["id"]])
        if item["id"] in lexical_rank:
            item["score"] += 1 / (60 + lexical_rank[item["id"]])
    return sorted(candidates, key=lambda c: c["score"], reverse=True)[:limit]


def citations_for(chunks):
    result = []
    for index, chunk in enumerate(chunks, 1):
        result.append({
            "number": index,
            "chunk_id": chunk["id"],
            "source_id": chunk["source_id"],
            "source_name": chunk["source_name"],
            "page": chunk["page"],
            "excerpt": chunk["text"][:700],
        })
    return result


def evidence_text(chunks):
    sections = []
    total = 0
    for index, chunk in enumerate(chunks, 1):
        location = f", page {chunk['page']}" if chunk["page"] else ""
        block = f"[SOURCE {index}: {chunk['source_name']}{location}]\n{chunk['text']}"
        if total + len(block) > 42_000:
            break
        sections.append(block)
        total += len(block)
    return "\n\n".join(sections)


SYSTEM_PROMPT = """You are Gemma Notebook, a source-grounded research assistant. Answer using only the supplied evidence. Every factual claim must carry a citation like [1] using the matching SOURCE number. Never invent citations. If the evidence is insufficient, say so plainly. Synthesize across sources, identify conflicts, and separate source claims from your analysis. Use clear Markdown."""


ARTIFACT_PROMPTS = {
    "summary": "Create a concise but comprehensive notebook overview: central topic, key claims, important evidence, disagreements, and open questions.",
    "faq": "Create a useful FAQ with 8-12 questions and source-grounded answers.",
    "guide": "Create a structured study and briefing guide with key concepts, evidence, terminology, review questions, and practical takeaways.",
}


async def ollama_stream(messages):
    payload = {
        "model": GENERATION_MODEL,
        "messages": messages,
        "stream": True,
        "options": {"num_ctx": 65536, "num_predict": 4096},
        "keep_alive": "15m",
    }
    async with httpx.AsyncClient(timeout=httpx.Timeout(600, connect=10)) as client:
        async with client.stream("POST", f"{OLLAMA_URL}/api/chat", json=payload) as response:
            response.raise_for_status()
            async for line in response.aiter_lines():
                if not line:
                    continue
                import json
                item = json.loads(line)
                if item.get("error"):
                    raise RuntimeError(item["error"])
                content = item.get("message", {}).get("content", "")
                if content:
                    yield content


def source_kind(filename):
    suffix = Path(filename).suffix.lower()
    kinds = {
        ".pdf": "pdf", ".docx": "docx", ".odt": "odt", ".txt": "text",
        ".md": "markdown", ".markdown": "markdown", ".html": "html", ".htm": "html", ".csv": "csv",
    }
    if suffix not in kinds:
        raise ValueError("Unsupported file type. Use PDF, DOCX, ODT, TXT, Markdown, HTML, or CSV.")
    return kinds[suffix]


def safe_filename(name):
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(name).name).strip(" .")
    return cleaned[:180] or "source.txt"
